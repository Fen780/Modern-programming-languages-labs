import os
import random
import time
from bs4 import BeautifulSoup
import openpyxl
import requests

session = requests.Session()
session.headers.update({
    'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36',
    'Accept': 'text/html,application/xhtml+xml,application/xml;q=0.9,image/avif,image/webp,image/apng,*/*;q=0.8,application/signed-exchange;v=b3;q=0.7',
    'Accept-Language': 'en-US,en;q=0.9,ru;q=0.8',
    'Connection': 'keep-alive',
    'Upgrade-Insecure-Requests': '1',
    'Sec-Fetch-Dest': 'document',
    'Sec-Fetch-Mode': 'navigate',
    'Sec-Fetch-Site': 'none',
    'Sec-Fetch-User': '?1',
})

def random_sleep():
    sleep_time = random.uniform(30, 60)
    print(f"Don't worry, juts sleep for {sleep_time} seconds")
    time.sleep(sleep_time)


def save_data(name, imo, mmsi, ship_type):
    filename='result.xlsx'

    if  os.path.exists(filename):
        workbook = openpyxl.load_workbook(filename)
        sheet = workbook.active
    else:
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.append(["Name", "IMO", "MMSI", "Type"])

    sheet.append([name, imo, mmsi, ship_type])

    workbook.save(filename)
    workbook.close()


def proceed_details(details_url):
    try:
        session.headers.update({'Referer': 'https://www.vesselfinder.com/'})
        random_sleep()
        resp = session.get(details_url, timeout=20)
        soup = BeautifulSoup(resp.text, 'html.parser')

        title_tag = soup.find('h1', class_='title')
        ship_name = title_tag.text.strip()

        imo = "-"
        mmsi = "-"
        ship_type = "-"

        rows = soup.find_all('tr')
        
        for row in rows:
            cells = row.find_all('td')
            if len(cells) != 2:
                continue
                
            key = cells[0].text.strip()
            value = cells[1].text.strip()

            if "IMO / MMSI" in key:
                parts = value.split('/')
                if len(parts) == 2:
                    imo = parts[0].strip()
                    mmsi = parts[1].strip()
                else:
                    mmsi = value
                    
            elif key == "MMSI":
                mmsi = value.strip()

            elif "AIS тип" in key:
                ship_type = value.strip()

        save_data(ship_name, imo, mmsi, ship_type)
    except Exception as e:
        print(f"Error reading url: {e}")
        return


def proceed_url(url):
    try:
        session.headers.update({'Referer': 'https://www.google.com/'})
        random_sleep()
        resp = session.get(url, timeout=20)
        soup = BeautifulSoup(resp.text, 'html.parser')
        ship_links = soup.find_all('a', class_='ship-link')
        count = len(ship_links)
        if count != 1:
            return
        link = ship_links[0]
        href = link.get('href')
        details_url = 'https://www.vesselfinder.com' + href
        proceed_details(details_url)
    except Exception as e:
        print(f"Error reading url: {e}")
        return

def proceed_sheet():
    file = 'Links.xlsx'
    try:
        workbook = openpyxl.load_workbook(file)
        sheet = workbook.active
    except FileNotFoundError:
        print('File not found')
        return
    
    counter = 0
    for row in sheet.iter_rows(min_row=2, min_col=1, max_col=1):
        cell = row[0]
        url = cell.value
        counter += 1
        if url is not None:
            print(f"Row {counter}: {url}")
            url = url.replace(" ", "%20")
            proceed_url(url)

def main():
    output_file = 'result.xlsx'
    if os.path.exists(output_file):
        os.remove(output_file)
    proceed_sheet()

if __name__ == "__main__":
    main()